"""Caption files for a batch of carousels: one .txt per deck plus a combined .txt and .docx.

Write captions in a JSON file (see examples/captions.json):

    {
      "hashtags": {"count": 5, "required": {"instagram": "#yourbrand", "linkedin": "#YourBrand"}},
      "decks": [
        {"slug": "lead-flow", "title": "Every enquiry makes the same journey",
         "instagram": "caption text ... #a #b #c #d #yourbrand",
         "linkedin":  "caption text ... #A #B #C #D #YourBrand"}
      ]
    }

    python -m rocket.captions examples/captions.json --out exports

Checks each caption has exactly `count` hashtags, includes the required tag, and warns when the
first line (the part shown before "more") is longer than 125 characters.
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

TAG = re.compile(r"(?<![\w&])#\w+")
PLATFORMS = ("instagram", "linkedin")


def validate(cfg: dict) -> list[str]:
    rules = cfg.get("hashtags", {})
    count = rules.get("count")
    required = rules.get("required", {})
    problems = []
    for d in cfg["decks"]:
        for p in PLATFORMS:
            text = d.get(p)
            if not text:
                problems.append(f"{d['slug']}: missing {p} caption")
                continue
            tags = TAG.findall(text)
            if count is not None and len(tags) != count:
                problems.append(f"{d['slug']} / {p}: {len(tags)} hashtags, expected {count}")
            need = required.get(p)
            if need and need.lower() not in (t.lower() for t in tags):
                problems.append(f"{d['slug']} / {p}: required tag {need} missing")
            hook = text.strip().splitlines()[0]
            if len(hook) > 125:
                print(f"  note: {d['slug']} / {p} first line is {len(hook)} chars; it may be cut before 'more'")
    return problems


def block(d: dict) -> str:
    return (f"{d.get('title', d['slug'])}\n{'=' * 60}\n\n"
            f"INSTAGRAM\n---------\n{d['instagram'].strip()}\n\n"
            f"LINKEDIN\n--------\n{d['linkedin'].strip()}\n")


def write(cfg: dict, out: pathlib.Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    for d in cfg["decks"]:
        (out / f"{d['slug']}-captions.txt").write_text(block(d), encoding="utf-8")
    combined = "\n\n".join(block(d) for d in cfg["decks"])
    (out / "ALL-CAPTIONS.txt").write_text(combined, encoding="utf-8")
    try:
        import docx
    except ImportError:
        print("  python-docx not installed; skipped ALL-CAPTIONS.docx  (pip install python-docx)")
        return
    doc = docx.Document()
    for i, d in enumerate(cfg["decks"]):
        if i:
            doc.add_page_break()
        doc.add_heading(d.get("title", d["slug"]), level=1)
        for p in PLATFORMS:
            doc.add_heading(p.capitalize(), level=2)
            for para in d[p].strip().split("\n"):
                doc.add_paragraph(para)
    doc.save(out / "ALL-CAPTIONS.docx")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", help="captions JSON")
    ap.add_argument("--out", default="exports", help="output folder")
    args = ap.parse_args(argv)
    cfg = json.loads(pathlib.Path(args.json).read_text(encoding="utf-8"))
    problems = validate(cfg)
    if problems:
        print("PROBLEMS:")
        for p in problems:
            print("  -", p)
        return 1
    write(cfg, pathlib.Path(args.out))
    print(f"wrote captions for {len(cfg['decks'])} decks to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
